# -*- coding: utf-8 -*-
# Parte 2 del deck (Eduardo Travi, AVI SPL): la prueba de concepto con SIP.
# Layouts genéricos y reutilizables. Todo el texto llega como datos desde
# content.py; aquí sólo vive la maquetación, con los mismos helpers y la
# misma paleta que la parte 1 para que el deck se lea como uno solo.
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
import diagrams as D
from diagrams import (rrect, oval, textbox, _settext, card, chevron, arrow, link,
                      INDIGO, BLUE, ORANGE, MAGENTA, TEAL, GREEN, RED, WHITE, GREY, LILAC)

LINE = RGBColor(0xDD, 0xDD, 0xE5)
LEFT = 0.92
FULL = 11.5


def _color(name):
    return {"blue": BLUE, "orange": ORANGE, "magenta": MAGENTA, "teal": TEAL,
            "green": GREEN, "red": RED, "indigo": INDIGO, "grey": GREY}.get(name, BLUE)


def _bullets(slide, x, y, w, h, items, size=11.5, color=INDIGO, anchor=MSO_ANCHOR.TOP, space=6):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    for j, it in enumerate(items):
        p = tf.paragraphs[0] if j == 0 else tf.add_paragraph(); p.space_after = Pt(space)
        r = p.add_run(); r.text = "•  " + it
        r.font.size = Pt(size); r.font.color.rgb = color; r.font.name = D.FONT
    return tb


def _banner(slide, y, text, h=0.85, size=13.5, sub=None):
    b = rrect(slide, LEFT, y, FULL, h, INDIGO, radius=0.1)
    items = [(text, size, True, WHITE, PP_ALIGN.CENTER)]
    if sub:
        items.append((sub, 10.5, False, LILAC, PP_ALIGN.CENTER))
    _settext(b.text_frame, items)
    return b


def _panel(slide, x, y, w, h, col, head, items, head_h=0.82, size=11.5, sub=None, head_shape="rect"):
    rrect(slide, x, y, w, h, WHITE, line=LINE, radius=0.06)
    if head_shape == "chevron":
        chevron(slide, x, y, w, head_h, col, head)
    else:
        hd = rrect(slide, x, y, w, head_h, col, radius=0.06)
        _settext(hd.text_frame, [(head, 13.5, True, WHITE, PP_ALIGN.CENTER)])
    y0 = y + head_h + 0.16
    if sub:
        textbox(slide, x + 0.28, y0, w - 0.56, 0.4, [(sub, 12, True, col, PP_ALIGN.LEFT)])
        y0 += 0.42
    _bullets(slide, x + 0.28, y0, w - 0.56, y + h - y0 - 0.12, items, size=size)


# ---------- layout: columnas de paneles (2 ó 3) + banner opcional ----------
def _twocol(slide, s):
    panels = s["panels"]
    n = len(panels); gap = 0.35 if n == 2 else 0.3
    w = (FULL - gap * (n - 1)) / n
    top = 1.8; has_banner = bool(s.get("banner"))
    h = 3.3 if has_banner else 4.35
    for i, p in enumerate(panels):
        x = LEFT + i * (w + gap)
        _panel(slide, x, top, w, h, _color(p["col"]), p["head"], p["items"],
               head_h=p.get("head_h", 0.82), size=p.get("size", 11.5), sub=p.get("sub"),
               head_shape=p.get("head_shape", "rect"))
    if has_banner:
        _banner(slide, top + h + 0.3, s["banner"], sub=s.get("banner_sub"))


# ---------- layout: lista de hosts a la izquierda + panel/comando a la derecha ----------
def _hosts(slide, s):
    hosts = s["hosts"]
    top = 1.95; hh = 0.72; gap = 0.14
    lw = 5.2
    col_h = len(hosts) * hh + (len(hosts) - 1) * gap
    for i, hst in enumerate(hosts):
        y = top + i * (hh + gap)
        col = _color(hst["col"])
        c = rrect(slide, LEFT, y, lw, hh, WHITE, line=col, radius=0.1)
        _settext(c.text_frame, [
            ("%s  (%s)" % (hst["name"], hst["ip"]), 12, True, col, PP_ALIGN.LEFT),
            (hst["role"], 9.5, False, GREY, PP_ALIGN.LEFT),
        ], anchor=MSO_ANCHOR.MIDDLE)
    rx = LEFT + lw + 0.6
    rw = 12.42 - rx
    if s.get("command"):
        # caja de comando (consola) + etiqueta
        lbl = s.get("command_label", "Comando")
        box = rrect(slide, rx, top, rw, col_h, INDIGO, radius=0.08)
        tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = lbl; r.font.size = Pt(12.5); r.font.bold = True
        r.font.color.rgb = LILAC; r.font.name = D.FONT
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER; p2.space_before = Pt(10)
        r2 = p2.add_run(); r2.text = s["command"]
        r2.font.size = Pt(13); r2.font.bold = True; r2.font.color.rgb = WHITE
        r2.font.name = "Consolas"
    else:
        rp = s["right"]
        col = _color(rp.get("col", "indigo"))
        rrect(slide, rx, top, rw, col_h, WHITE, line=LINE, radius=0.08)
        rrect(slide, rx, top, rw, 0.7, col, radius=0.08)
        textbox(slide, rx, top, rw, 0.7, [(rp["head"], 13.5, True, WHITE, PP_ALIGN.CENTER)],
                anchor=MSO_ANCHOR.MIDDLE)
        if rp.get("quote"):
            q = rrect(slide, rx + 0.3, top + 0.95, rw - 0.6, col_h - 1.3, D.PALE, radius=0.08)
            _settext(q.text_frame, [(rp["quote"], 13, False, INDIGO, PP_ALIGN.CENTER)])
        else:
            _bullets(slide, rx + 0.3, top + 0.9, rw - 0.6, col_h - 1.1, rp["items"], size=11.5)
    if s.get("banner"):
        _banner(slide, top + col_h + 0.3, s["banner"], sub=s.get("banner_sub"))


# ---------- layout: mapeo "falla → control" (remediación) ----------
def _mapping(slide, s):
    rows = s["rows"]
    top = 1.85; h = 0.6; gap = 0.13
    lw = 5.1; aw = 0.5; rw = FULL - lw - aw - 0.3
    for i, (left, right) in enumerate(rows):
        y = top + i * (h + gap)
        col = D.PALETTE[i % len(D.PALETTE)]
        lb = rrect(slide, LEFT, y, lw, h, WHITE, line=col, radius=0.1)
        _settext(lb.text_frame, [(left, 11.5, True, col, PP_ALIGN.LEFT)])
        arrow(slide, LEFT + lw + 0.05, y + h / 2, LEFT + lw + aw - 0.05, y + h / 2, color=col, width=2.0)
        rb = rrect(slide, LEFT + lw + aw, y, rw, h, WHITE, line=LINE, radius=0.1)
        _settext(rb.text_frame, [(right, 11.5, False, INDIGO, PP_ALIGN.LEFT)])
    if s.get("banner"):
        bottom = top + len(rows) * h + (len(rows) - 1) * gap
        _banner(slide, bottom + 0.2, s["banner"], h=0.72, sub=s.get("banner_sub"))


# ---------- layout: fila de tarjetas (reutiliza el helper de la parte 1) ----------
def _cards(slide, prs, s):
    cards = [(_color(c["col"]), c["head"], c["desc"]) for c in s["cards"]]
    top = 2.1 if not s.get("subtitle") else 2.5
    if s.get("subtitle"):
        textbox(slide, LEFT, 1.75, FULL, 0.5, [(s["subtitle"], 13, False, GREY, PP_ALIGN.CENTER)],
                anchor=MSO_ANCHOR.MIDDLE)
    D.cards_row(slide, prs, cards, top=top, h=3.0)
    if s.get("banner"):
        _banner(slide, top + 3.0 + 0.3, s["banner"], sub=s.get("banner_sub"))


_LAYOUTS = {"twocol": _twocol, "hosts": _hosts, "mapping": _mapping}


def d_poc(slide, prs, s):
    if s.get("tag"):
        textbox(slide, LEFT, 1.55, FULL, 0.38, [(s["tag"], 12.5, True, BLUE, PP_ALIGN.LEFT)],
                anchor=MSO_ANCHOR.MIDDLE)
    lay = s["layout"]
    if lay == "cards":
        _cards(slide, prs, s)
    elif lay in ("twocol", "mapping"):
        _LAYOUTS[lay](slide, s)
    elif lay == "hosts":
        _hosts(slide, s)
