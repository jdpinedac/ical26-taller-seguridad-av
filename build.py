# -*- coding: utf-8 -*-
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image
import zipfile, shutil, re, os, subprocess
import content as C
from embed_fonts import embed as embed_fonts
import diagrams

FONT = "Montserrat"
FONT_MED = "Montserrat Medium"

TPL = "template.pptx"
OUT = "ICAL26-seguridad-AV-parte1.pptx"
PYXIS_W = "assets/pyxis-blanco.png"   # claro, para fondos oscuros/degradado
PYXIS_D = "assets/pyxis.png"          # color, para fondo blanco
LOCKUP  = "assets/logos/ical-lockup.png"  # ic26+infocomm recortado
AVISPL_W = "assets/logos/avispl-blanco.png"  # AVI SPL blanco (co-presentador)
AVISPL_D = "assets/logos/avispl-oscuro.png"  # AVI SPL indigo (fondo blanco)

INDIGO = RGBColor(0x34, 0x1E, 0x66)
BLUE   = RGBColor(0x00, 0x7E, 0xE5)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
GREY   = RGBColor(0x6B, 0x6B, 0x75)
LILAC  = RGBColor(0xD6, 0xCF, 0xEA)

L_COVER, L_CONTENT, L_DARK, L_DARKSP = 0, 2, 7, 3

def delete_all_slides(prs):
    part = prs.part
    for sid in list(prs.slides._sldIdLst):
        rid = sid.get(qn('r:id'))
        prs.slides._sldIdLst.remove(sid)
        if rid in part.rels:
            part.drop_rel(rid)

def ph_by_idx(slide, idx):
    for ph in slide.placeholders:
        if ph.placeholder_format.idx == idx:
            return ph
    return None

def set_title(slide, text):
    ph = ph_by_idx(slide, 0)
    if ph is not None:
        ph.text = text
        for r in ph.text_frame.paragraphs[0].runs:
            r.font.name = FONT_MED
    return ph

def no_bullet(p):
    pPr = p._p.get_or_add_pPr()
    for tag in ('a:buChar', 'a:buAutoNum'):
        e = pPr.find(qn(tag))
        if e is not None:
            pPr.remove(e)
    pPr.append(pPr.makeelement(qn('a:buNone'), {}))

def fill_bullets(tf, bullets, color=None, size=None, shrink=True):
    tf.word_wrap = True
    if shrink:
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    first = True
    for (txt, lvl, bold) in bullets:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = lvl
        r = p.add_run(); r.text = txt; r.font.bold = bold; r.font.name = FONT
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = color

def pic_height(path, width_in):
    iw, ih = Image.open(path).size
    return Inches(width_in * ih / iw)

def add_logo(slide, prs, path, corner):
    w_in = 1.5
    h = pic_height(path, w_in); w = Inches(w_in)
    if corner == "br":
        left = prs.slide_width - w - Inches(0.45)
        top = prs.slide_height - h - Inches(0.28)
    else:  # bl
        left = Inches(0.45)
        top = prs.slide_height - h - Inches(0.28)
    slide.shapes.add_picture(path, left, top, width=w)

def add_textbox(slide, l, t, w, h, lines, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    first = True
    for (txt, size, bold, color) in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.alignment = align
        r = p.add_run(); r.text = txt
        r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color; r.font.name = FONT
    return tb

def add_shot_box(slide, l, t, w, h, caption):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(l), Inches(t), Inches(w), Inches(h))
    box.shadow.inherit = False
    box.fill.solid(); box.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF5)
    box.line.color.rgb = BLUE; box.line.width = Pt(1.25)
    ln = box.line._get_or_add_ln()
    ln.append(ln.makeelement(qn('a:prstDash'), {'val': 'dash'}))
    tf = box.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = "[ " + caption + " ]"
    r.font.size = Pt(13); r.font.italic = True; r.font.color.rgb = GREY; r.font.name = FONT
    return box

def dark_statement(slide, big, small):
    obj = ph_by_idx(slide, 1)
    tf = obj.text_frame; tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    first = True
    for ln in big.split("\n"):
        p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
        p.alignment = PP_ALIGN.CENTER; no_bullet(p)
        r = p.add_run(); r.text = ln
        r.font.size = Pt(32); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
    p = tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER; no_bullet(p); p.space_before = Pt(18)
    r = p.add_run(); r.text = small
    r.font.size = Pt(18); r.font.color.rgb = LILAC; r.font.name = FONT


def swap_fonts(path):
    tmp = path + ".tmp"
    repl = {
        "Gotham Medium": FONT_MED,
        "Gotham Book": FONT,
        "Gotham": FONT,
        "Arial Rounded MT Bold": FONT_MED,
        "Aptos Display": FONT_MED,
        "Aptos": FONT,
    }
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename.endswith(".xml") and (
                item.filename.startswith("ppt/theme/")
                or item.filename.startswith("ppt/slideMasters/")
                or item.filename.startswith("ppt/slideLayouts/")
                or item.filename.startswith("ppt/slides/")
            ):
                t = data.decode("utf-8")
                for a, b in repl.items():
                    t = t.replace('typeface="%s"' % a, 'typeface="%s"' % b)
                data = t.encode("utf-8")
            zout.writestr(item, data)
    os.replace(tmp, path)


def photo_bg(slide, prs, path):
    from PIL import Image as _Img
    iw, ih = _Img.open(path).size
    sw, sh = prs.slide_width, prs.slide_height
    # llenar: si al ajustar por ancho sobra alto, uso ancho; si no, uso alto
    disp_h = int(sw * ih / iw)
    if disp_h >= sh:
        pic = slide.shapes.add_picture(path, 0, int((sh - disp_h) / 2), width=sw)
    else:
        disp_w = int(sh * iw / ih)
        pic = slide.shapes.add_picture(path, int((sw - disp_w) / 2), 0, height=sh)
    ov = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, sw, sh)
    ov.shadow.inherit = False
    ov.fill.solid(); ov.fill.fore_color.rgb = RGBColor(0x18, 0x10, 0x2C)
    srgb = ov.fill._xPr.find(qn('a:solidFill')).find(qn('a:srgbClr'))
    srgb.append(srgb.makeelement(qn('a:alpha'), {'val': '62000'}))
    ov.line.fill.background()
    tree = slide.shapes._spTree
    for el in (pic._element, ov._element):
        tree.remove(el)
    tree.insert(2, pic._element)
    tree.insert(3, ov._element)
    # remontar lockup ic26+infocomm abajo a la derecha
    lw = Inches(1.95); lh = pic_height(LOCKUP, 1.95)
    slide.shapes.add_picture(LOCKUP, sw - lw - Inches(0.45), sh - lh - Inches(0.3), width=lw)



def add_mock(slide, prs, path, caption, y_min=3.7, y_max=6.5, max_w=10.2, max_h=2.5):
    from PIL import Image as _I
    iw, ih = _I.open(path).size
    w = max_w; h = w * ih / iw
    if h > max_h:
        h = max_h; w = h * iw / ih
    left = (13.333 - w) / 2
    cap_h = 0.32
    top = y_min + ((y_max - y_min) - (h + cap_h)) / 2   # centrado vertical en la franja libre
    pic = slide.shapes.add_picture(path, Inches(left), Inches(top), width=Inches(w))
    pic.line.color.rgb = RGBColor(0xC9, 0xCD, 0xD3); pic.line.width = Pt(0.75)
    add_textbox(slide, 0.92, top + h + 0.04, 11.5, 0.32,
                [("Ejemplo representativo: " + caption, 9.5, False, GREY)], align=PP_ALIGN.CENTER)


def main():
    prs = Presentation(TPL)
    delete_all_slides(prs)
    L = prs.slide_layouts

    for s in C.SLIDES:
        k = s["kind"]

        if k == "cover":
            slide = prs.slides.add_slide(L[L_COVER])
            # dejar vacios los placeholders del template (el hero infocomm es el centro)
            # bloque de titulo en el tercio inferior, blanco, sobre el degradado
            add_textbox(slide, 0.0, 4.55, 13.333, 0.9,
                        [(s["title"], 32, True, WHITE)], align=PP_ALIGN.CENTER)
            add_textbox(slide, 0.0, 5.35, 13.333, 0.5,
                        [(s["subtitle"], 18, False, WHITE)], align=PP_ALIGN.CENTER)
            add_textbox(slide, 0.0, 5.86, 13.333, 0.4,
                        [(s["fecha_lugar"], 13, True, LILAC)], align=PP_ALIGN.CENTER)
            add_textbox(slide, 0.0, 6.18, 13.333, 0.8,
                        [(s["presenter1"], 12.5, True, WHITE),
                         (s["presenter2"], 12.5, True, WHITE)], align=PP_ALIGN.CENTER)
            add_logo(slide, prs, PYXIS_W, "bl")
            # AVI SPL (co-presentador) abajo a la derecha
            aw = Inches(1.9); ah = pic_height(AVISPL_W, 1.9)
            slide.shapes.add_picture(AVISPL_W, prs.slide_width - aw - Inches(0.45),
                                     prs.slide_height - ah - Inches(0.34), width=aw)

        elif k == "content":
            slide = prs.slides.add_slide(L[L_CONTENT])
            set_title(slide, s["title"])
            if s.get("diagram"):
                diagrams.REGISTRY[s["diagram"]](slide, prs)
            else:
                fill_bullets(ph_by_idx(slide, 10).text_frame, s["bullets"])
                if s.get("shot"):
                    add_shot_box(slide, 1.73, 4.5, 9.6, 1.75, s["shot"])

        elif k == "demo":
            slide = prs.slides.add_slide(L[L_CONTENT])
            set_title(slide, s["title"])
            if s.get("logo"):
                lw = Inches(1.15); lh = pic_height(s["logo"], 1.15)
                slide.shapes.add_picture(s["logo"], prs.slide_width - lw - Inches(0.5), Inches(0.5), width=lw)
            if s.get("badge"):
                bsh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, prs.slide_width - Inches(2.3), Inches(0.6), Inches(1.8), Inches(0.5))
                bsh.shadow.inherit = False; bsh.adjustments[0] = 0.5
                bsh.fill.solid(); bsh.fill.fore_color.rgb = RGBColor(0x1A,0x6B,0xB5); bsh.line.fill.background()
                p = bsh.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
                r = p.add_run(); r.text = s["badge"]; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = WHITE; r.font.name = FONT
            tb = slide.shapes.add_textbox(Inches(1.73), Inches(1.7), Inches(9.9), Inches(2.3))
            tf = tb.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            r = p.add_run(); r.text = s["tag"]; r.font.size = Pt(13); r.font.bold = True; r.font.color.rgb = BLUE; r.font.name = FONT
            p.space_after = Pt(6)
            for (txt, lvl, bold) in s["bullets"]:
                p = tf.add_paragraph(); p.space_after = Pt(8)
                r = p.add_run(); r.text = txt; r.font.size = Pt(15); r.font.color.rgb = INDIGO; r.font.name = FONT
            if s.get("mock"):
                add_mock(slide, prs, s["mock"], s["shot"])
            else:
                add_shot_box(slide, 1.73, 4.2, 9.6, 2.0, s["shot"])

        elif k == "statement":
            slide = prs.slides.add_slide(L[L_DARK])
            if s.get("photo"):
                photo_bg(slide, prs, s["photo"])
            dark_statement(slide, s["big"], s["small"])
            add_logo(slide, prs, PYXIS_W, "bl")

        elif k == "section":
            slide = prs.slides.add_slide(L[L_DARKSP])
            dark_statement(slide, s["big"], s["small"])

        elif k == "closing":
            slide = prs.slides.add_slide(L[L_CONTENT])
            set_title(slide, s["title"])
            diagrams.d_kit(slide, prs, s.get("shot", "QR al kit descargable"))

    prs.save(OUT)
    swap_fonts(OUT)
    embed_fonts(OUT)
    # exportar PDF; publicar a la carpeta indicada por ICAL26_PUBLISH_DIR (si está definida)
    pdf = OUT.replace(".pptx", ".pdf")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", ".", OUT],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=200)
    dest = os.environ.get("ICAL26_PUBLISH_DIR")
    n = len(prs.slides._sldIdLst)
    if dest:
        os.makedirs(dest, exist_ok=True)
        for f in (OUT, pdf):
            shutil.copy(f, os.path.join(dest, os.path.basename(f)))
        print("Guardado", OUT, "+ PDF, publicados en", dest, "(", n, "slides )")
    else:
        print("Guardado", OUT, "+ PDF local (", n, "slides ). Define ICAL26_PUBLISH_DIR para publicar.")

if __name__ == "__main__":
    main()
