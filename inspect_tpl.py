from pptx import Presentation
from pptx.util import Emu
p = Presentation("template.pptx")
print("slide size", p.slide_width, p.slide_height)
for i, lay in enumerate(p.slide_layouts):
    print(f"\n== layout[{i}] name={lay.name!r}")
    for ph in lay.placeholders:
        f = ph.placeholder_format
        print(f"   ph idx={f.idx} type={f.type} name={ph.name!r} pos=({Emu(ph.left).inches:.2f},{Emu(ph.top).inches:.2f}) size=({Emu(ph.width).inches:.2f}x{Emu(ph.height).inches:.2f})")
    for sh in lay.shapes:
        if not sh.is_placeholder:
            print(f"   shape {sh.shape_type} name={sh.name!r} pos=({Emu(sh.left).inches:.2f},{Emu(sh.top).inches:.2f}) size=({Emu(sh.width).inches:.2f}x{Emu(sh.height).inches:.2f})")
print("\n== slides in template:")
for i, s in enumerate(p.slides):
    print(i, s.slide_layout.name, [sh.name for sh in s.shapes])
